#!/usr/bin/env python3
"""Erzeugt apps/server/src/data/bls.json aus dem Bundeslebensmittelschlüssel (BLS) des Max Rubner-Instituts.

Der BLS steht unter CC BY 4.0 frei zum Download (https://blsdb.de/download). Übernommen werden je
Lebensmittel Code, Name und die Werte der üblichen Nährwerttabelle je 100 g.

Aufruf im Repo-Root, mit der Excel-Datei aus dem Download:

    python3 scripts/create-bls-data.py /pfad/zu/BLS_4_0_Daten_2025_DE.xlsx

Braucht openpyxl (`pip install openpyxl`).
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'apps/server/src/data/bls.json'

SOURCE = (
    'Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), Version 4.0 – Deutsche Nährstoffdatenbank. '
    'Karlsruhe. DOI: 10.25826/Data20251217-134202-0'
)

# Nährstoff in der App → Code im BLS (Spalten heißen z. B. „ENERCC Energie (Kilokalorien) [kcal/100g]“)
NUTRIENTS = {
    'kcal': 'ENERCC',
    'fat': 'FAT',
    'saturatedFat': 'FASAT',
    'carbs': 'CHO',
    'sugar': 'SUGAR',
    'fiber': 'FIBT',
    'protein': 'PROT625',
    'salt': 'NACL',
}


def rounded(value, digits):
    if value is None or value == '':
        return None
    if isinstance(value, str):
        # „<LOD“ und Ähnliches: unter der Nachweisgrenze, also praktisch nichts
        if value.strip().startswith('<'):
            return 0
        try:
            value = float(value.replace(',', '.'))
        except ValueError:
            return None
    number = float(value)
    return round(number) if digits == 0 else round(number, digits)


def main(path: str) -> None:
    sheet = load_workbook(path, read_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {}
    for key, code in NUTRIENTS.items():
        matches = [i for i, title in enumerate(header) if isinstance(title, str) and title.startswith(f'{code} ') and '[' in title]
        if len(matches) != 1:
            raise SystemExit(f'Spalte für {code} nicht eindeutig gefunden')
        columns[key] = matches[0]

    foods = []
    for row in rows:
        code, name = row[0], row[1]
        if not code or not name:
            continue
        values = [rounded(row[columns[key]], 0 if key == 'kcal' else 1) for key in NUTRIENTS]
        foods.append([code, ' '.join(str(name).split()), *values])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    data = {'source': SOURCE, 'license': 'CC BY 4.0', 'nutrients': list(NUTRIENTS), 'foods': foods}
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    missing = sum(1 for food in foods if any(value is None for value in food[2:]))
    print(f'{len(foods)} Lebensmittel nach {OUT.relative_to(ROOT)}, {missing} mit Lücken, {OUT.stat().st_size // 1024} KB')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
